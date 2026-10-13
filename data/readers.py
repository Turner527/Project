"""Чтение отчётов: xlsx (все листы, форматы чисел, заметки) и csv (кодировка и разделитель определяются сами)."""
import csv
import io
import re
from dataclasses import dataclass, field

import pandas as pd

from data.parsing import is_empty
from data.validation import WorkbookError, check_zip_size

MAX_ROWS = 600_000
MAX_NOTES = 40
HEADER_SCAN_ROWS = 30
NOTE_SHEET_RE = re.compile(r"о файле|описани|справк|readme|notes?|info|легенд|инструкц|about", re.I)


@dataclass
class RawFile:
    filename: str
    kind: str  # xlsx | csv
    frames: dict[str, pd.DataFrame]  # листы-таблицы: значения как есть, первая строка данных = индекс 0
    header_rows: dict[str, int]  # номер строки заголовка в файле (с 1) по листам
    formats: dict[str, dict[str, str]]  # лист -> столбец -> формат числа Excel первой строки данных
    notes: list[str] = field(default_factory=list)
    sheet: str = ""  # выбранный лист с данными

    @property
    def sheets(self) -> list[str]:
        return list(self.frames)


def _find_header(rows: list[tuple]) -> int:
    """Индекс строки заголовка среди первых строк: ≥2 непустых текстовых ячеек, как можно ближе к началу."""
    counts = [sum(not is_empty(c) for c in r) for r in rows]
    widest = max(counts, default=0)
    for i, r in enumerate(rows):
        cells = [c for c in r if not is_empty(c)]
        strings = [c for c in cells if isinstance(c, str)]
        if (len(cells) >= 2 and len(cells) >= 0.6 * widest and len(strings) >= 0.8 * len(cells)
                and len({str(c).strip().lower() for c in strings}) == len(strings)):
            return i
    return 0


def _frame(rows: list[tuple], header_idx: int) -> pd.DataFrame | None:
    header = rows[header_idx]
    width = max((len(r) for r in rows[header_idx:]), default=0)
    names, seen = [], {}
    for i in range(width):
        raw = header[i] if i < len(header) else None
        name = re.sub(r"\s+", " ", str(raw)).strip() if not is_empty(raw) else f"Столбец {i + 1}"
        seen[name] = seen.get(name, 0) + 1
        names.append(name if seen[name] == 1 else f"{name} ({seen[name]})")
    body = [list(r) + [None] * (width - len(r)) for r in rows[header_idx + 1:]]
    df = pd.DataFrame(body, columns=names, dtype=object)
    df = df.dropna(how="all")
    keep = [c for c in df.columns if not (c.startswith("Столбец ") and df[c].map(is_empty).all())]
    df = df[keep].reset_index(drop=True)
    return df if len(df) and df.shape[1] >= 2 else None


def _notes(rows: list[tuple]) -> list[str]:
    out = []
    for r in rows:
        for c in r:
            if isinstance(c, str) and not c.strip().startswith("=") and len(c.strip()) > 3:
                out.append(re.sub(r"\s+", " ", c).strip()[:240])
    return out


def _read_xlsx(data: bytes) -> RawFile:
    import openpyxl

    check_zip_size(data)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise WorkbookError([
            "Не удалось открыть файл как Excel (.xlsx): файл повреждён, защищён паролем или имеет другой формат "
            f"({type(exc).__name__}). Сохраните книгу как .xlsx без пароля и загрузите снова."
        ]) from exc
    frames, headers, formats, notes = {}, {}, {}, []
    for ws in wb.worksheets:
        rows = []
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            rows.append(r)
            if i > MAX_ROWS:
                raise WorkbookError([f"На листе «{ws.title}» больше {MAX_ROWS // 1000} тысяч строк — файл слишком большой."])
        while rows and all(is_empty(c) for c in rows[-1]):
            rows.pop()
        if not rows:
            continue
        note_sheet = bool(NOTE_SHEET_RE.search(ws.title))
        if note_sheet:
            notes += _notes(rows)  # листы «О файле» / «Описание»: показываем пользователю, модели не передаём
        h = _find_header(rows[:HEADER_SCAN_ROWS])
        df = _frame(rows, h)
        if df is None or len(df) < 3:
            if not note_sheet:
                notes += _notes(rows)
            continue
        frames[ws.title], headers[ws.title] = df, h + 1
        if not note_sheet:
            notes += _notes(rows[:h])  # заголовки отчёта над таблицей
        first = next(iter(ws.iter_rows(min_row=h + 2, max_row=h + 2)), ())
        formats[ws.title] = {name: getattr(cell, "number_format", "General") or "General"
                             for name, cell in zip(df.columns, first)}
    wb.close()
    if not frames:
        raise WorkbookError(["В файле не найдено ни одной таблицы: нужен лист с заголовками столбцов и строками данных."])
    return RawFile("", "xlsx", frames, headers, formats, notes[:MAX_NOTES])


def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1251", errors="replace")


def _read_csv(data: bytes, filename: str) -> RawFile:
    text = _decode(data)
    sample = text[:8192]
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
    rows: list[tuple] = []
    for i, r in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter)):
        rows.append(tuple(r))
        if i > MAX_ROWS:
            raise WorkbookError([f"В файле больше {MAX_ROWS // 1000} тысяч строк — файл слишком большой."])
    while rows and all(is_empty(c) for c in rows[-1]):
        rows.pop()
    header_idx = _find_header(rows[:HEADER_SCAN_ROWS]) if rows else 0
    df = _frame(rows, header_idx) if rows else None
    if df is None or len(df) < 3:
        raise WorkbookError(["В csv-файле слишком мало данных: нужны строка заголовков, минимум два столбца и три строки данных."])
    name = re.sub(r"\.csv$", "", filename, flags=re.I) or "Данные"
    return RawFile("", "csv", {name: df}, {name: header_idx + 1}, {name: {}}, [])


def read_file(data: bytes, filename: str) -> RawFile:
    """Читает xlsx или csv и выбирает лист с данными. Бросает WorkbookError с понятным сообщением."""
    lower = filename.lower()
    if lower.endswith(".csv"):
        raw = _read_csv(data, filename)
    elif lower.endswith(".xlsx") or data.startswith(b"PK\x03\x04"):
        raw = _read_xlsx(data)
    elif lower.endswith(".xls"):
        raise WorkbookError(["Формат .xls устарел: откройте файл в Excel и сохраните как .xlsx или .csv."])
    else:
        raise WorkbookError([f"Файл «{filename}» не в поддерживаемом формате. Загрузите .xlsx или .csv."])
    raw.filename = filename
    raw.sheet = pick_sheet(raw)
    return raw


def pick_sheet(raw: RawFile) -> str:
    """Лист с данными: больше строк × столбцов; листы с «описанием» и подобным уступают."""
    def score(name: str) -> float:
        df = raw.frames[name]
        return len(df) * df.shape[1] * (0.05 if NOTE_SHEET_RE.search(name) else 1)

    return max(raw.frames, key=score)
